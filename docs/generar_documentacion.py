from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Documentacion_estado_actual_plataforma.docx"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tc_pr.append(shd)

def borders(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr(); b = tc_pr.first_child_found_in('w:tcBorders')
    if b is None: b = OxmlElement('w:tcBorders'); tc_pr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge; e = b.find(qn(tag))
        if e is None: e = OxmlElement(tag); b.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),color)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(3)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(9.5)
    if color: r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)

def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,'17365D'); set_cell_text(c,h,True,(255,255,255))
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            if n % 2 == 1: shade(cells[i],'F3F7FB')
            set_cell_text(cells[i],str(v))
    if widths:
        for row in t.rows:
            for cell,w in zip(row.cells,widths): cell.width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

def heading(doc, text, level=1):
    p=doc.add_heading(text, level=level); p.paragraph_format.space_before=Pt(14 if level==1 else 9); p.paragraph_format.space_after=Pt(6)
    for r in p.runs: r.font.color.rgb=RGBColor(0,0,0)
    return p

def para(doc, text, bold_prefix=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(7); p.paragraph_format.line_spacing=1.12
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold=True; p.add_run(text[len(bold_prefix):])
    else: p.add_run(text)
    return p

def bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.add_run(item)

def figure_placeholder(doc, number, title, filename, purpose):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(f"[INSERTAR DIAGRAMA {number}: {title}]"); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(31,78,121)
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(8)
    q.add_run(f"Archivo D2: docs/diagramas-d2/{filename}. {purpose}").italic=True

def remove_paragraph_border(paragraph_or_style):
    ppr = paragraph_or_style._element.get_or_add_pPr()
    border = ppr.find(qn('w:pBdr'))
    if border is not None:
        ppr.remove(border)

doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos'); styles['Normal'].font.size=Pt(10.5)
for name,size in [('Title',25),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
    styles[name].font.name='Aptos Display' if name!='Normal' else 'Aptos'; styles[name].font.size=Pt(size); styles[name].font.color.rgb=RGBColor(0,0,0)
remove_paragraph_border(styles['Title'])
header=sec.header.paragraphs[0]; header.text='TECHNICAL ASSESSMENT PLATFORM  |  DOCUMENTACIÓN TÉCNICA'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor(89,89,89)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('Estado documentado: 20 de septiembre de 2026  |  Página ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Documentación del estado actual de la plataforma')
remove_paragraph_border(p)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Proceso constructivo, decisiones de diseño y guía de diagramas'); r.italic=True; r.font.size=Pt(13)
doc.add_paragraph('\n')
add_table(doc,['Elemento','Alcance documentado'],[
 ['Producto','Plataforma de evaluaciones técnicas de programación'],
 ['Estado','Implementación actual del repositorio'],
 ['Arquitectura','Monolito modular: React/Vite, NestJS, PostgreSQL/Prisma y Judge0'],
 ['Propósito','Explicar cómo se construyó la solución y por qué se tomaron las decisiones principales'],
 ['Diagramas','Seis fuentes D2 independientes, listas para compilar e insertar'],
],[1.6,4.9])
para(doc,'Este documento describe el estado actual de la Technical Assessment Platform a partir de su implementación. La solución separa la experiencia de candidato y la administración, mantiene reglas de negocio en una API NestJS modular y delega la ejecución de código no confiable a Judge0. El resultado es una base extensible para asignar retos, controlar intentos y evaluar soluciones de programación sin ejecutar código arbitrario dentro del backend.')

heading(doc,'1 Alcance y lectura del documento')
para(doc,'La documentación está orientada a explicar decisiones verificables en el código actual, no a reconstruir decisiones históricas que no quedaron registradas. El orden constructivo expuesto es una lectura técnica razonada: parte del modelo de dominio y llega a la experiencia, la seguridad y el despliegue.')
bullets(doc,['Los marcadores de diagrama indican el lugar exacto para insertar cada imagen compilada desde D2.','Los archivos fuente están en docs/diagramas-d2 y pueden exportarse a SVG o PNG.','User y AssessmentAttempt son las fuentes únicas de identidad del candidato y del resultado de la evaluación.'])

heading(doc,'2 Entendimiento del problema')
para(doc,'El problema se entendió como una plataforma para administrar y resolver evaluaciones técnicas de programación. Un administrador configura evaluaciones, preguntas, lenguajes permitidos, casos de prueba y ventanas de disponibilidad. Un candidato autenticado solo consulta las evaluaciones que le fueron asignadas, inicia un intento con vencimiento y recibe resultados de sus ejecuciones.')
heading(doc,'Actores y objetivos',2)
add_table(doc,['Actor','Objetivo principal','Límites relevantes'],[
 ['Administrador','Crear, publicar, asignar y revisar evaluaciones.','Solo puede asignar evaluaciones publicadas y debe definir una ventana válida.'],
 ['Candidato','Resolver ejercicios y consultar su resultado.','Solo accede a evaluaciones asignadas, dentro de disponibilidad y con JWT.'],
 ['Judge0','Compilar y ejecutar código.','Es externo a la API; recibe límites de recursos y red deshabilitada.'],
],[1.25,2.8,2.45])
figure_placeholder(doc,1,'casos de uso','01-casos-de-uso.d2','Resume las responsabilidades de candidato y administrador.')

heading(doc,'3 Proceso constructivo y decisiones de diseño')
heading(doc,'3.1 Modelar el núcleo antes de la interfaz',2)
para(doc,'El primer bloque debía responder qué se evalúa, quién lo resuelve y cómo se conserva evidencia. Por eso el modelo Prisma define Assessment, Question, TestCase, Submission y TestResult, junto con usuarios, asignaciones e intentos. Las relaciones, restricciones únicas e índices expresan invariantes directamente en la persistencia: la posición de una pregunta es única dentro de una evaluación y un resultado de prueba es único para una submission y un caso.')
para(doc,'Decisión de diseño: se almacenan los casos de prueba y los resultados por separado. Esto permite que una misma pregunta tenga casos visibles y ocultos, que cada ejecución sea auditable y que el detalle privado no se devuelva al candidato.')
figure_placeholder(doc,2,'modelo de dominio y datos','02-modelo-dominio-y-datos.d2','Representa entidades principales, atributos de negocio y cardinalidades.')
heading(doc,'3.2 Organizar la API como monolito modular',2)
para(doc,'Con el dominio definido, el backend se dividió por módulos funcionales: autenticación, asignaciones, evaluaciones, preguntas, intentos, submissions, ejecución y evaluación. Cada módulo separa presentación, aplicación, dominio e infraestructura. Los controladores se mantienen como entrada HTTP; los casos de uso contienen la coordinación y las reglas; Prisma implementa repositorios; Judge0 queda detrás de un puerto de dominio.')
para(doc,'Decisión de diseño: se eligió un monolito modular en lugar de microservicios. Para una plataforma de este tamaño reduce complejidad operacional y mantiene límites internos claros. El puerto CodeExecutionPort deja abierta la sustitución de Judge0 sin propagar dependencia del proveedor por la aplicación.')
figure_placeholder(doc,3,'componentes de la solución','03-componentes.d2','Muestra fronteras entre frontend, módulos de API, persistencia y ejecución externa.')
heading(doc,'3.3 Construir el flujo de evaluación con control de tiempo',2)
para(doc,'El candidato inicia un AssessmentAttempt solo si la evaluación está publicada y tiene una asignación vigente. El servidor calcula expiresAt y vuelve a validar que el intento siga activo antes y después de cada llamada al proveedor. Si hay un intento activo se reutiliza; si ya se completó, se devuelve el completado. Así, la regla de tiempo no depende del reloj del navegador.')
para(doc,'Decisión de diseño: el progreso se recalcula desde las submissions evaluadas del intento. El resultado de cada pregunta considera su última submission evaluada y el puntaje agregado se pondera con el score configurado para cada pregunta. Cuando todas están evaluadas el intento se completa; de otro modo se actualiza su resumen.')
heading(doc,'3.4 Ejecutar código no confiable fuera del backend',2)
para(doc,'La ejecución reclama primero la submission para evitar dobles ejecuciones concurrentes. Después itera los casos de prueba: envía código, lenguaje y entrada a Judge0, normaliza la respuesta, compara la salida y persiste un TestResult. Al finalizar, calcula el puntaje y actualiza el intento. Ante un error se libera el reclamo para que la submission no quede bloqueada.')
para(doc,'Decisión de diseño: la API nunca ejecuta código del candidato localmente. Judge0 recibe límites configurables de CPU, tiempo de pared, memoria, procesos y tamaño de archivo; además se solicita la red deshabilitada. Los casos ocultos se persisten y participan en el puntaje, pero su entrada, salida esperada, detalle de ejecución y mensaje se omiten de la respuesta del candidato.')
figure_placeholder(doc,4,'secuencia de ejecución de una submission','04-secuencia-ejecucion-submission.d2','Explica la secuencia desde el editor hasta el reporte de resultados.')
heading(doc,'3.5 Cerrar los límites de acceso y capacidad',2)
para(doc,'La capa de autenticación usa JWT y roles ADMIN/CANDIDATE. El registro público crea únicamente candidatos. En la API, la asignación, publicación, disponibilidad, propiedad del intento, lenguajes permitidos, duración y cuotas se validan en servidor. La documentación del repositorio establece un máximo de 30 000 caracteres por submission y 20 submissions por intento.')
para(doc,'Decisión de diseño: la protección no se limita a ocultar pantallas. Las rutas frontend se protegen para la experiencia y los guards/validaciones backend protegen el dato. Esta doble capa evita que una llamada directa a la API permita consultar o ejecutar recursos fuera de autorización.')
figure_placeholder(doc,5,'estados del intento y de la submission','05-estados-del-intento-y-submission.d2','Sintetiza las transiciones que condicionan ejecución y cierre.')
heading(doc,'3.6 Preparar la operación local y el despliegue',2)
para(doc,'Para desarrollo local se usa Docker Compose únicamente para PostgreSQL; el monorepo pnpm permite ejecutar API y web en paralelo. Para producción, Terraform define una alternativa de coste contenido en AWS: frontend estático en S3 privado detrás de CloudFront, API NestJS en EC2 con Nginx y PostgreSQL en RDS privado. CloudFront enruta /api hacia la API, por lo que el navegador consume una ruta relativa y evita una URL de backend expuesta.')
para(doc,'Decisión de diseño: la instancia EC2 pública puede salir a Judge0 sin un NAT Gateway, mientras RDS permanece en una red privada y solo acepta tráfico de la instancia. Esta configuración favorece bajo coste para una demostración; no pretende proporcionar alta disponibilidad ni un Judge0 privado.')
figure_placeholder(doc,6,'despliegue objetivo en AWS','06-despliegue-aws.d2','Muestra la topología descrita por la infraestructura Terraform.')

heading(doc,'4 Arquitectura actual')
add_table(doc,['Capa','Tecnología','Responsabilidad'],[
 ['Cliente','React, Vite y Monaco Editor','Rutas por rol, editor, temporizador y consumo de API.'],
 ['API','NestJS y TypeScript','Reglas de negocio, JWT, validación, coordinación de casos de uso.'],
 ['Persistencia','PostgreSQL y Prisma','Modelo relacional, migraciones y repositorios.'],
 ['Ejecución','Judge0 mediante adaptador','Sandbox de compilación y ejecución de código no confiable.'],
 ['Infraestructura','Docker Compose y Terraform','Entorno local y despliegue AWS de bajo coste.'],
],[1.1,1.75,3.65])
heading(doc,'Módulos de negocio',2)
bullets(doc,['Auth: registro, login, perfil, hash de contraseña, JWT y roles.','Assignment: asignación de evaluaciones a candidatos y ventana de disponibilidad.','Assessment y Question: ciclo de vida de evaluaciones, preguntas, lenguajes y casos de prueba.','AssessmentAttempt: inicio, vigencia, resumen y finalización del intento actual.','Submission, Execution y Evaluation: ejecución aislada, persistencia de resultados y puntuación.'])

heading(doc,'5 Decisiones de diseño resumidas')
add_table(doc,['Decisión','Motivación','Consecuencia'],[
 ['Monolito modular','Separar responsabilidades sin distribuir la operación prematuramente.','Módulos sustituibles internamente con despliegue simple.'],
 ['Prisma con PostgreSQL','Relaciones y restricciones del dominio requieren consistencia.','Migraciones versionadas y consultas centralizadas en repositorios.'],
 ['Intento en servidor','El tiempo y autorización no deben depender del navegador.','Vigencia comprobada en operaciones sensibles.'],
 ['Puerto de ejecución','El proveedor de sandbox debe ser intercambiable.','Judge0 queda aislado en un adaptador de infraestructura.'],
 ['Casos ocultos','Evitar revelar la solución de validación.','El candidato recibe solo información permitida.'],
 ['CloudFront con ruta relativa API','Simplificar consumo seguro del frontend.','Sin mixed content ni endpoint directo configurado en navegador.'],
],[1.25,2.65,2.6])

heading(doc,'6 Guía para compilar e insertar los diagramas')
para(doc,'Las fuentes se entregan separadas para que puedan mantenerse con el código. Desde la raíz del repositorio, con D2 instalado, se puede generar cada imagen con el siguiente patrón:')
p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.35); r=p.add_run('d2 docs/diagramas-d2/01-casos-de-uso.d2 docs/diagramas-d2/01-casos-de-uso.svg'); r.font.name='Consolas'; r.font.size=Pt(9)
para(doc,'Compile cada archivo a SVG o PNG e inserte el resultado en el marcador con el mismo número. SVG suele conservar mejor la nitidez; PNG es una alternativa si Word o su destino de publicación lo requieren. Los seis archivos se listan también en docs/diagramas-d2/README.md.')
add_table(doc,['Marcador Word','Fuente D2','Vista'],[
 ['Diagrama 1','01-casos-de-uso.d2','Casos de uso'],['Diagrama 2','02-modelo-dominio-y-datos.d2','Modelo de dominio y datos'],['Diagrama 3','03-componentes.d2','Componentes'],['Diagrama 4','04-secuencia-ejecucion-submission.d2','Secuencia'],['Diagrama 5','05-estados-del-intento-y-submission.d2','Estados'],['Diagrama 6','06-despliegue-aws.d2','Despliegue'],
],[1.45,3.2,1.85])

heading(doc,'7 Límites actuales y siguientes mejoras')
para(doc,'La infraestructura de producción está planteada como demostración de bajo coste y no incluye Multi-AZ, balanceador, RDS Proxy, NAT Gateway ni una instancia privada de Judge0. El README también advierte que las credenciales de seed son exclusivas de desarrollo y que la cuenta administrativa de producción debe crearse con un mecanismo privado y un JWT_SECRET robusto.')
bullets(doc,['Definir auditoría de eventos y métricas operativas para ejecuciones, errores del proveedor y tiempos de respuesta.','Evaluar un Judge0 privado o proveedor equivalente si aumentan requisitos de aislamiento, capacidad o disponibilidad.','Agregar pruebas automatizadas de dominio y contratos además de los scripts actuales de verificación de endpoints y ejecución.'])

heading(doc,'Anexo A Evidencia consultada')
para(doc,'La documentación se elaboró a partir de README.md, apps/api/prisma/schema.prisma, módulos y casos de uso de NestJS, rutas React/Vite, docker-compose.yml e infraestructura/terraform. Los nombres de las entidades, límites y flujos citados reflejan ese estado de implementación.')
doc.save(OUT)
print(OUT)
