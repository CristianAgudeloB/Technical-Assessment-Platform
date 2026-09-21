from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'Guion_presentacion_construccion_plataforma.docx'
DIAGRAMS = 'docs/diagramas-presentacion-d2'

def remove_paragraph_border(target):
    ppr = target._element.get_or_add_pPr()
    border = ppr.find(qn('w:pBdr'))
    if border is not None:
        ppr.remove(border)

def set_page(section, landscape=False):
    section.top_margin = Inches(.62)
    section.bottom_margin = Inches(.62)
    section.left_margin = Inches(.82)
    section.right_margin = Inches(.82)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)

def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = 'TECHNICAL ASSESSMENT PLATFORM  |  GUION DE PRESENTACIÓN'
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(89, 89, 89)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Guía para explicar la construcción del programa  |  Página ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    p.add_run(text)
    return p

def bullets(doc, values):
    for value in values:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(value)

def placeholder(doc, number, name, filename, note, height=150):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f'[INSERTAR DIAGRAMA {number}: {name.upper()}]')
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(height)
    r = p.add_run('Espacio reservado para el diagrama')
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(89, 89, 89)

def simple_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ''
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell._tc.get_or_add_tcPr().append(_shade('17365D'))
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ''
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                cells[index]._tc.get_or_add_tcPr().append(_shade('F3F7FB'))
            run = cells[index].paragraphs[0].add_run(value)
            run.font.size = Pt(9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def _shade(fill):
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), fill)
    return shade

doc = Document()
set_page(doc.sections[0])
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
styles['Normal'].font.size = Pt(10.8)
for name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12)]:
    styles[name].font.name = 'Aptos Display'
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
remove_paragraph_border(styles['Title'])

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run('Construcción de la plataforma de evaluaciones técnicas')
remove_paragraph_border(title)
para(doc, 'La plataforma permite crear, asignar y resolver evaluaciones técnicas de programación. Su construcción se centró en controlar el acceso de cada candidato, evaluar código de forma segura y conservar resultados verificables.')
heading(doc, '1 Entendimiento del problema')
para(doc, 'La plataforma debía permitir administrar pruebas técnicas de programación y resolverlas en un entorno controlado. El administrador prepara y asigna retos; el candidato accede solo a los que tiene disponibles, los resuelve y consulta su resultado.')
heading(doc, 'Decisiones que guiaron esta etapa', 2)
bullets(doc, [
    'Se separaron los roles de administrador y candidato para que cada persona vea únicamente las acciones que necesita.',
    'Se incluyó la asignación con fechas de disponibilidad para controlar quién puede presentar cada evaluación y cuándo.',
    'Se planteó la evaluación como un proceso de preguntas, intentos, envíos y resultados para conservar evidencia del recorrido.',
])
placeholder(doc, 1, 'entendimiento del problema', '01-entendimiento-del-problema.d2', '', 160)

db_section = doc.add_section(WD_SECTION.NEW_PAGE)
set_page(db_section, landscape=True)
heading(doc, '2 La base de datos como punto de partida')
para(doc, 'Antes de construir pantallas y endpoints, se definió qué información debía persistir y cómo se relacionaría. El modelo permite conservar la configuración de la evaluación, la asignación del candidato, el intento realizado y la evidencia de cada ejecución.')
heading(doc, 'Por qué se tomaron estas decisiones', 2)
bullets(doc, [
    'Las entidades se separaron para mantener trazabilidad entre evaluación, pregunta, caso de prueba, envío y resultado.',
    'Los casos ocultos permiten validar soluciones sin revelar la lógica completa de evaluación.',
    'Las relaciones y reglas únicas reducen duplicados y protegen la consistencia de posiciones, asignaciones y resultados.',
])
placeholder(doc, 2, 'modelo de datos', '02-modelo-de-datos.d2', '', 220)

portrait = doc.add_section(WD_SECTION.NEW_PAGE)
set_page(portrait)
heading(doc, '3 Arquitectura de la solución')
para(doc, 'La solución se construyó como un monolito modular: una aplicación web para los usuarios, una API que concentra las reglas del negocio, una base de datos relacional y un servicio externo que ejecuta el código de forma aislada.')
heading(doc, 'Por qué se tomaron estas decisiones', 2)
bullets(doc, [
    'React y Vite permiten una interfaz rápida y separada para la experiencia de candidato y administración.',
    'NestJS organiza las reglas por módulos y evita mezclar autenticación, evaluaciones, intentos y ejecución.',
    'Judge0 se usa fuera del backend para no ejecutar código no confiable en el servidor de la aplicación.',
])
placeholder(doc, 3, 'arquitectura de la solución', '03-arquitectura-de-la-solucion.d2', '', 150)

doc.add_page_break()
heading(doc, '4 Recorrido de una evaluación')
para(doc, 'El candidato inicia sesión, consulta una evaluación asignada y comienza un intento que tiene hora de vencimiento. Después resuelve las preguntas, envía soluciones y recibe un resultado calculado con los casos de prueba configurados.')
heading(doc, 'Por qué se tomaron estas decisiones', 2)
bullets(doc, [
    'El tiempo se valida en el servidor para que no dependa del reloj del navegador.',
    'Un intento activo puede recuperarse para evitar que el candidato pierda su progreso al recargar la aplicación.',
    'La plataforma guarda el avance y completa el intento cuando todas las preguntas tienen resultado.',
])
placeholder(doc, 4, 'actividad del candidato', '04-actividad-del-candidato.d2', '', 150)

doc.add_page_break()
heading(doc, '5 Ejecución y resultados')
para(doc, 'Cada solución se valida contra los casos de prueba de la pregunta. La API verifica el intento, delega la ejecución a Judge0, compara el resultado, guarda la evidencia y actualiza el progreso de la evaluación.')
heading(doc, 'Por qué se tomaron estas decisiones', 2)
bullets(doc, [
    'La submission se reclama antes de ejecutarse para evitar dos evaluaciones concurrentes del mismo envío.',
    'Cada resultado se guarda por caso de prueba para poder explicar y auditar el puntaje obtenido.',
    'La respuesta al candidato oculta la información de los casos confidenciales, aunque estos sí cuenten en la calificación.',
])
placeholder(doc, 5, 'secuencia de ejecución', '05-secuencia-de-ejecucion.d2', '', 145)

doc.add_page_break()
heading(doc, '6 Estados y controles')
para(doc, 'Los estados hacen visible cuándo un intento o una solución puede avanzar. La autenticación, los roles, las asignaciones, los límites de ejecución y el vencimiento se validan en la API, no solo en la interfaz.')
heading(doc, 'Por qué se tomaron estas decisiones', 2)
bullets(doc, [
    'Los estados evitan ejecutar o modificar una solución que ya fue evaluada o cuyo intento venció.',
    'JWT y roles protegen los recursos de candidato y administrador en la experiencia y en el backend.',
    'Los límites de CPU, memoria, procesos, tiempo y red reducen el riesgo al evaluar código de terceros.',
])
placeholder(doc, 6, 'estados principales', '06-estados-principales.d2', '', 145)

doc.add_page_break()
heading(doc, '7 Operación y despliegue')
para(doc, 'Para desarrollo se usa PostgreSQL en Docker Compose y se ejecutan frontend y API en paralelo. Para producción, Terraform plantea un frontend estático detrás de CloudFront, una API en EC2 y PostgreSQL en RDS privado.')
heading(doc, 'Por qué se tomaron estas decisiones', 2)
bullets(doc, [
    'Docker Compose simplifica levantar la dependencia local de base de datos.',
    'CloudFront y S3 sirven la aplicación web de forma eficiente, mientras la ruta /api dirige las solicitudes al backend.',
    'RDS permanece privado y la API es el único componente que debe comunicarse con la base de datos.',
])
placeholder(doc, 7, 'despliegue', '07-despliegue.d2', '', 155)
heading(doc, 'Conclusión', 2)
para(doc, 'La plataforma se construyó separando las responsabilidades que más importan: administrar evaluaciones, controlar el acceso y el tiempo, ejecutar código de forma segura y conservar resultados verificables.')

doc.save(OUT)
print(OUT)
