from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / 'Guion_presentacion_construccion_plataforma.docx'
OUT = ROOT / 'docs' / 'Presentacion_con_controles_de_seguridad.docx'

def add_paragraph(document, text, italic=False, centered=False, size=None, color=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    if centered:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return paragraph

def add_bullet(document, text):
    paragraph = document.add_paragraph(style='List Bullet')
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)
    return paragraph

document = Document(SOURCE)
document.add_page_break()

document.add_heading('7 Seguridad y ejecución aislada', level=1)
add_paragraph(
    document,
    'La ejecución de código se diseñó para que una solución escrita por un candidato no se ejecute de forma nativa dentro de la API ni del equipo que ejecuta la plataforma.',
)

document.add_heading('Aislamiento de la ejecución', level=2)
add_bullet(document, 'La API envía el código a Judge0 mediante CodeExecutionPort; el backend solo coordina la solicitud, recibe el resultado y lo persiste.')
add_bullet(document, 'Judge0 ejecuta el código en un entorno externo y aislado con red deshabilitada, por lo que la solución no puede acceder libremente a servicios internos.')
add_bullet(document, 'Cada ejecución tiene límites configurables: 2 segundos de CPU, 5 segundos de tiempo de pared, 128 MB de memoria, 32 procesos y 1 MB de archivos por defecto.')
add_bullet(document, 'La API también impone un límite total de espera para el proveedor; si no hay respuesta, la ejecución se controla como un error y no queda abierta indefinidamente.')

document.add_heading('Validación y control de acceso', level=2)
add_bullet(document, 'Las rutas usan JWT y roles. Además, el servidor valida que la evaluación esté asignada, esté disponible y que el intento permanezca activo antes y después de ejecutar código.')
add_bullet(document, 'El código recibido se valida antes de procesarse: máximo 30 000 caracteres, lenguaje permitido y hasta 20 submissions por intento.')
add_bullet(document, 'Los casos de prueba ocultos participan en la calificación, pero su entrada, salida esperada y detalle de ejecución no se devuelven al candidato.')

document.add_heading('Docker y manejo de configuración', level=2)
add_bullet(document, 'Docker Compose levanta PostgreSQL para un entorno local reproducible. La ejecución de soluciones no ocurre en ese contenedor ni en el contenedor de la API.')
add_bullet(document, 'La imagen de la API se ejecuta con un usuario de aplicación no privilegiado. Esto reduce el alcance del proceso incluso dentro de su propio contenedor.')
add_bullet(document, 'Las variables sensibles se mantienen fuera del repositorio mediante .env y .gitignore; el archivo .env.example contiene únicamente valores simulados.')

placeholder = document.add_paragraph()
placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
placeholder.paragraph_format.space_before = Pt(10)
placeholder.paragraph_format.space_after = Pt(3)
run = placeholder.add_run('[INSERTAR DIAGRAMA: SEGURIDAD Y EJECUCIÓN AISLADA]')
run.bold = True
run.font.size = Pt(10.5)
run.font.color.rgb = RGBColor(0, 0, 0)

space = document.add_paragraph()
space.alignment = WD_ALIGN_PARAGRAPH.CENTER
space.paragraph_format.space_after = Pt(140)
run = space.add_run('Espacio reservado para el diagrama')
run.italic = True
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(89, 89, 89)

document.save(OUT)
print(OUT)
